"""
엑셀 -> 아이라벨 붙여넣기용 변환기 (Streamlit 웹앱 버전)
------------------------------------------------------
엑셀 파일을 업로드하면, 함수(수식) 결과값만 뽑아서 한 줄에 하나씩
정리해서 보여줍니다. 화면에 나온 복사 버튼으로 복사한 뒤 아이라벨에
붙여넣으면 됩니다.

실행 방법:
    pip install -r requirements.txt
    streamlit run main.py
"""

import streamlit as st
import openpyxl
from io import BytesIO

st.set_page_config(page_title="엑셀 → 아이라벨 변환기", page_icon="🏷️")

st.title("🏷️ 엑셀 → 아이라벨 변환기")
st.write("엑셀 파일의 함수 결과값만 뽑아서, 아이라벨에 한 번에 붙여넣을 수 있는 형태로 만들어줍니다.")

uploaded_file = st.file_uploader("엑셀 파일 업로드 (.xlsx)", type=["xlsx"])

if uploaded_file is not None:
    # data_only=True: 수식이 아니라 엑셀이 마지막으로 계산해둔 결과값을 읽음
    wb = openpyxl.load_workbook(BytesIO(uploaded_file.getvalue()), data_only=True)

    sheet_name = st.selectbox("시트 선택", wb.sheetnames)
    ws = wb[sheet_name]

    col1, col2, col3 = st.columns(3)
    with col1:
        col_letter = st.text_input("값이 있는 열 (예: A, B, C)", value="A").strip().upper()
    with col2:
        start_row = st.number_input("시작 행", min_value=1, value=1, step=1)
    with col3:
        end_row = st.number_input("끝 행", min_value=1, value=ws.max_row, step=1)

    if st.button("값 추출하기", type="primary"):
        values = []
        for row in range(int(start_row), int(end_row) + 1):
            cell = ws[f"{col_letter}{row}"]
            if cell.value is not None:
                values.append(str(cell.value))

        if not values:
            st.warning("선택한 범위에서 값을 찾지 못했어요. 열/행 범위를 확인해보세요.")
        else:
            st.success(f"{len(values)}개의 값을 찾았어요. 아래 오른쪽 위 복사 버튼으로 복사한 뒤 아이라벨에 붙여넣으세요.")
            text = "\n".join(values)
            st.code(text, language=None)

            st.download_button(
                label="텍스트 파일로 다운로드",
                data=text,
                file_name="아이라벨_붙여넣기용.txt",
                mime="text/plain",
            )

st.divider()
st.caption(
    "엑셀에 수식이 들어있는 셀의 계산된 값을 읽으려면, 그 파일을 Excel에서 한 번 열어서 "
    "저장(Ctrl+S)한 상태여야 합니다. 그래야 계산 결과가 파일 안에 저장되어 있어요."
)
